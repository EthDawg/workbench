#!/usr/bin/env python3
"""Synthetic regressions for the bundled, user-supplied PowerPoint helpers.

Run with Python plus the bundled skill's requirements.txt dependencies.
No real session, installed app, network, or slide-rendering service is used.
"""
import copy
import hashlib
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from pptx import Presentation

PACK = Path(__file__).resolve().parents[1] / "Sources/LocalVoice/Resources/build-snap-and-talk-deck/packs/servicenow-employee-experience/1.0.0"
sys.dont_write_bytecode = True
sys.path.insert(0, str(PACK / "scripts"))
import build_deck
import session_io
import session_outline


class DeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="snap-talk-deck-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.session = self.root / "Session with spaces"
        self.session.mkdir()
        sections = []
        # Short/repeated text and non-16:9 captures must NOT be discarded.
        for name, narration in [("a", "  Um, hello.\n\n"), ("b", "  Um, hello.\n\n")]:
            directory = self.session / "items" / name
            directory.mkdir(parents=True)
            shutil.copyfile(PACK / "brand/assets/logo_wordmark.png", directory / "screen.png")
            (directory / "narration.txt").write_bytes(narration.encode())
            sections.append({"id": name, "directory": f"items/{name}", "status": "ready",
                             "screenshot": f"items/{name}/screen.png", "transcript": f"items/{name}/narration.txt"})
        sections += [{"id": "c", "directory": "trash/c", "status": "ready"},
                     {"id": "d", "directory": "items/d", "status": "queued"},
                     {"id": "e", "directory": "items/e", "status": "ready", "screenshot": "items/e/missing.png"}]
        self.manifest = {"formatVersion": 1, "title": "Synthetic walkthrough", "sections": sections}
        self.save_manifest()
        self.draft = self.session / "outline.json"
        session_outline.main(str(self.session), str(self.draft))
        self.outline = json.loads(self.draft.read_bytes())
        for slide in self.outline["chapters"][0]["slides"]:
            slide.update(headline="A greeting", takeaways=["The narrator says hello"])

    def save_manifest(self):
        (self.session / "session.json").write_text(json.dumps(self.manifest))

    def build(self, outline=None, destination=None):
        self.draft.write_text(json.dumps(outline or self.outline))
        output = destination or self.session / "deck.pptx"
        build_deck.main(str(self.draft), str(self.session), str(output))
        return output

    def test_exact_notes_order_images_and_sources(self):
        originals = {p: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
                     for p in self.session.rglob("*") if p.is_file() and p != self.draft}
        deck = Presentation(self.build())
        self.assertEqual(len(deck.slides), 2)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=3)
        for index, slide in enumerate(deck.slides):
            self.assertEqual(slide.notes_slide.notes_text_frame.text, "  Um, hello.\n\n")
            picture = next(s for s in slide.shapes if s.name == "Screenshot")
            self.assertEqual(picture.image.blob, (self.session / f"items/{'ab'[index]}/screen.png").read_bytes())
            self.assertEqual((picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom), (0, 0, 0, 0))
            self.assertAlmostEqual(picture.width / picture.height, picture.image.size[0] / picture.image.size[1], places=4)
            self.assertTrue(any("The narrator says hello" in s.text for s in slide.shapes if s.has_text_frame))
        for path, original in originals.items():
            self.assertEqual((hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns), original)
        self.assertEqual(len(self.outline["excluded"]), 3)

    def test_short_repeated_and_wide_captures_kept(self):
        slides = self.outline["chapters"][0]["slides"]
        self.assertEqual([s["section_id"] for s in slides], ["a", "b"])
        self.assertTrue(all(s["flags"] for s in slides))

    def test_reject_changes_to_order_notes_and_images(self):
        for key in ("section_id", "notes", "image"):
            with self.subTest(key=key):
                outline = copy.deepcopy(self.outline)
                outline["chapters"][0]["slides"][0][key] = "changed"
                with self.assertRaises(ValueError): self.build(outline)
        self.outline["chapters"][0]["slides"].reverse()
        with self.assertRaises(ValueError): self.build()
        self.assertFalse((self.session / "deck.pptx").exists())

    def test_reject_overflow_instead_of_truncating(self):
        for values in ({"headline": "x" * 61}, {"takeaways": ["x"] * 4}, {"takeaways": ["x" * 61]}):
            outline = copy.deepcopy(self.outline)
            outline["chapters"][0]["slides"][0].update(values)
            with self.assertRaises(ValueError): self.build(outline)

    def test_refuse_overwrite_outline_deck_or_source(self):
        with self.assertRaises(FileExistsError): session_outline.main(str(self.session), str(self.draft))
        output = self.build()
        before = output.read_bytes()
        with self.assertRaises(FileExistsError): self.build()
        self.assertEqual(output.read_bytes(), before)
        with self.assertRaises(FileExistsError): self.build(destination=self.session / "session.json")

    def test_paths_and_symlinks_rejected(self):
        for path in ("/etc/passwd", "../outside", "items/a/../../outside", "items//a", "items\\a"):
            with self.subTest(path=path), self.assertRaises(ValueError):
                session_io.local_file(self.session, path)
        (self.session / "items/a/link.txt").symlink_to(self.session / "items/a/narration.txt")
        with self.assertRaises(ValueError): session_io.local_file(self.session, "items/a/link.txt")
        self.manifest["sections"][0]["transcript"] = "items/b/narration.txt"
        self.save_manifest()
        _, slides, excluded = session_io.captures(self.session)
        self.assertEqual([s["section_id"] for s in slides], ["b"])
        self.assertIn("own section", excluded[0]["reason"])

    def test_explicit_extra_slides(self):
        self.outline.update(cover=True, dividers=True, closing=True, summary={
            "eyebrow": "Summary", "title_white": "Greeting", "title_green": " overview",
            "cards": [{"head": "Greeting", "body": "The narrator says hello"}] * 4})
        deck = Presentation(self.build())
        self.assertEqual(len(deck.slides), 6)
        self.assertEqual(deck.slides[2].notes_slide.notes_text_frame.text, "  Um, hello.\n\n")

    def test_unsupported_format(self):
        self.manifest["formatVersion"] = 999
        self.save_manifest()
        with self.assertRaises(ValueError): self.build()


if __name__ == "__main__":
    unittest.main()
