#!/usr/bin/env python3
"""Build a brand-aligned Snap & Talk deck from an outline JSON - no template file needed.

usage: python build_deck.py outline.json "<session folder>" out.pptx

Everything visual comes from ../brand/brand.json and ../brand/assets/.
Requires: python-pptx, Pillow, lxml.
"""
import io, json, os, sys, datetime
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from session_io import captures, local_file, new_output

HERE = os.path.dirname(os.path.abspath(__file__))
BRAND_DIR = os.path.join(HERE, "..", "brand")
B = json.load(open(os.path.join(BRAND_DIR, "brand.json")))
C, F = B["colours"], B["fonts"]
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

def col(name_or_hex):
    return C.get(name_or_hex, name_or_hex)

def asset(key):
    return os.path.join(BRAND_DIR, B["assets"][key])

# ---------------------------------------------------------------- primitives
def box(slide, g):
    return Inches(g["x"]), Inches(g["y"]), Inches(g["w"]), Inches(g["h"])

def picture(slide, path, g):
    return slide.shapes.add_picture(path, *box(slide, g))

def background(slide, key):
    slide.shapes.add_picture(asset(key), 0, 0, Inches(B["canvas"]["width"]), Inches(B["canvas"]["height"]))

def text(slide, g, paras, anchor="t", insets=(0, 0, 0, 0), name=None):
    """paras: list of dicts {runs:[{t, font, size, bold, colour, tracking}], line (pct float) | line_pts, space_after, align}"""
    tb = slide.shapes.add_textbox(*box(slide, g))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    l, t, r, b = insets
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = Inches(l), Inches(t), Inches(r), Inches(b)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "ctr": PP_ALIGN.CENTER}[p.get("align", "l")]
        if "line_pts" in p: para.line_spacing = Pt(p["line_pts"])
        elif "line" in p: para.line_spacing = p["line"]
        para.space_after = Pt(p.get("space_after", 0))
        para.space_before = Pt(0)
        for rd in p["runs"]:
            run = para.add_run()
            run.text = rd["t"]
            f = run.font
            f.name = F[rd.get("font", "body")]
            f.size = Pt(rd["size"])
            f.bold = rd.get("bold", False)
            f.color.rgb = RGBColor.from_string(col(rd.get("colour", "white")))
            if rd.get("tracking"):
                run._r.get_or_add_rPr().set("spc", str(rd["tracking"]))
    # PowerPoint must not shrink/grow our boxes
    bp = tf._txBody.find(A + "bodyPr")
    for t_ in ("spAutoFit", "normAutofit", "noAutofit"):
        for e in bp.findall(A + t_): bp.remove(e)
    etree.SubElement(bp, A + "noAutofit")
    return tb

def _alpha_fill(parent_tag, hexcol, alpha):
    fill = etree.Element(A + parent_tag)
    c = etree.SubElement(fill, A + "srgbClr", val=hexcol)
    etree.SubElement(c, A + "alpha", val=str(alpha))
    return fill

def _shadow(spPr, s):
    eff = etree.SubElement(spPr, A + "effectLst")
    sh = etree.SubElement(eff, A + "outerShdw", blurRad=str(s["blurRad"]), dist=str(s["dist"]), dir=str(s["dir"]), algn="t", rotWithShape="0")
    c = etree.SubElement(sh, A + "srgbClr", val=col(s["colour"]))
    etree.SubElement(c, A + "alpha", val=str(s["alpha"]))

def card(slide, g):
    spec = B["card"]
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *box(slide, g))
    shp.adjustments[0] = spec["radius_adj"] / 100000
    spPr = shp._element.spPr
    for tag in ("solidFill", "gradFill", "noFill", "ln", "effectLst"):
        for e in spPr.findall(A + tag): spPr.remove(e)
    grad = etree.SubElement(spPr, A + "gradFill", flip="none", rotWithShape="1")
    gs = etree.SubElement(grad, A + "gsLst")
    for st in spec["fill_stops"]:
        g_ = etree.SubElement(gs, A + "gs", pos=str(st["pos"]))
        c = etree.SubElement(g_, A + "srgbClr", val=col(st["colour"]))
        etree.SubElement(c, A + "alpha", val=str(st["alpha"]))
    path = etree.SubElement(grad, A + "path", path="rect")
    etree.SubElement(path, A + "fillToRect", l="50000", t="50000", r="50000", b="50000")
    etree.SubElement(grad, A + "tileRect")
    ln = etree.SubElement(spPr, A + "ln", w=str(int(spec["line"]["width_pt"] * 12700)))
    ln.append(_alpha_fill("solidFill", col(spec["line"]["colour"]), spec["line"]["alpha"]))
    _shadow(spPr, spec["shadow"])
    shp.text_frame.text = ""
    shp.name = "Card"
    return shp

def rule(slide, x, y, length=None):
    L = length or B["rule"]["length"]
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + L), Inches(y))
    ln.line.color.rgb = RGBColor.from_string(col(B["rule"]["colour"]))
    ln.line.width = Pt(B["rule"]["width_pt"])
    ln.name = "Rule"
    return ln

def framed_screenshot(slide, img_path, alt):
    s = B["slides"]["content"]["screenshot"]
    with Image.open(img_path) as im:
        ar = im.width / im.height
    # letterbox inside the 16:9 box - never stretch
    bw, bh = s["w"], s["h"]
    w, h = (bw, bw / ar) if ar >= bw / bh else (bh * ar, bh)
    x, y = s["x"] + (bw - w) / 2, s["y"] + (bh - h) / 2
    # Keep original image bytes and every corner. Rounded masks clip screen UI.
    pic = slide.shapes.add_picture(str(img_path), Inches(x), Inches(y), Inches(w), Inches(h))
    spPr = pic._element.spPr
    ln = etree.SubElement(spPr, A + "ln", w=str(int(s["line"]["width_pt"] * 12700)))
    ln.append(_alpha_fill("solidFill", col(s["line"]["colour"]), s["line"]["alpha"]))
    _shadow(spPr, s["shadow"])
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.name = "Screenshot"
    return pic

def footer(slide, n, year):
    f = B["footer"]
    picture(slide, asset("logo"), f["logo"])
    text(slide, f["copyright"], [{"runs": [{"t": f["copyright"]["text"].format(year=year), "font": f["copyright"].get("font", "body"),
                                              "size": f["copyright"]["size"]}], "align": "r"}], anchor="ctr", insets=(0.1, 0.05, 0.1, 0.05))
    tb = text(slide, f["number"], [{"runs": [{"t": str(n), "size": f["number"]["size"]}], "align": {"left": "l", "right": "r"}[f["number"].get("align", "right")]}], anchor="ctr")
    # turn the typed number into a live slide-number field so reordering keeps it right
    r = tb.text_frame.paragraphs[0].runs[0]._r
    fld = etree.Element(A + "fld", id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    for child in list(r): fld.append(child)
    r.addprevious(fld); r.getparent().remove(r)

def notes(slide, s):
    if s: slide.notes_slide.notes_text_frame.text = s

# ---------------------------------------------------------------- slide types
def cover(prs, o, year):
    S = B["slides"]["cover"]; sl = prs.slides.add_slide(BLANK)
    background(sl, S["background"])
    picture(sl, asset("decoration"), S["decoration"])
    picture(sl, asset("logo"), S["logo"])
    t = S["title"]
    text(sl, t, [{"runs": [{"t": o["title"]["lead"] + " ", "font": "display", "size": t["size"], "bold": True, "colour": "green"},
                          {"t": o["title"].get("rest", ""), "font": "display", "size": t["size"], "bold": True, "colour": "white"}],
                 "line": t["line"]}], anchor=t["anchor"])
    st = S["subtitle"]
    text(sl, st, [{"runs": [{"t": o.get("subtitle", ""), "font": "medium", "size": st["size"]}], "line": st["line"]}],
         insets=(0.1, 0.05, 0.1, 0.05))
    pr = S["presenter"]
    text(sl, pr, [{"runs": [{"t": o.get("presenter", ""), "font": "body", "size": pr["size"], "bold": pr.get("bold", False), "colour": "green"}], "line": pr["line"]}], anchor="ctr")
    ro = S["role"]
    text(sl, ro, [{"runs": [{"t": o.get("role", ""), "font": "body", "size": ro["size"]}], "line": ro["line"]}],
         insets=(0.1, 0.05, 0.1, 0.05))
    footer(sl, len(prs.slides), year)
    notes(sl, o.get("cover_notes"))

def divider(prs, n, name, note):
    S = B["slides"]["divider"]; sl = prs.slides.add_slide(BLANK)
    background(sl, S["background"])
    picture(sl, asset("decoration"), S["decoration"])
    picture(sl, asset("logo"), S["logo"])
    t = S["title"]
    text(sl, t, [{"runs": [{"t": f"{n:02d}", "font": "display", "size": t["size"], "bold": True, "colour": "green"}], "line": t["line"]},
                 {"runs": [{"t": name, "font": "display", "size": t["size"], "bold": True, "colour": "white"}], "line": t["line"]}],
         anchor=t["anchor"])
    notes(sl, note)

def content(prs, i, chapter, s, session_dir, year):
    S = B["slides"]["content"]; sl = prs.slides.add_slide(BLANK)
    background(sl, S["background"])
    card(sl, S["card"])
    head = s["headline"]
    size = next(r["size"] for r in S["headline"]["size_rule"] if len(head) <= r["max_chars"]) if len(head) <= 60 else S["headline"]["size_rule"][-1]["size"]
    text(sl, S["headline_box"], [
        {"runs": [{"t": f"{i:02d}", "size": S["number"]["size"], "bold": True, "colour": "white", "tracking": -50},
                  {"t": "   " + chapter.upper(), "size": S["eyebrow"]["size"], "bold": True, "colour": "green"}],
         "line_pts": S["number"]["line_pts"]},
        {"runs": [{"t": head, "size": size, "bold": True, "colour": "green", "tracking": S["headline"]["tracking"]}],
         "line": S["headline"]["line"]}], anchor="b", insets=(0.1, 0.05, 0.1, 0.05))
    rule(sl, S["rule"]["x"], S["rule"]["y"])
    tk = S["takeaway"]
    text(sl, S["body_box"], [{"runs": [{"t": p, "size": tk["size"]}], "line": 1.0, "space_after": tk["space_after_pt"]}
                            for p in s["takeaways"]], insets=(0.1, 0.05, 0.1, 0.05))
    framed_screenshot(sl, local_file(session_dir, s["image"]), head)
    footer(sl, len(prs.slides), year)
    notes(sl, s.get("notes"))

def summary(prs, o, year):
    S = B["slides"]["summary"]; sl = prs.slides.add_slide(BLANK)
    background(sl, S["background"])
    text(sl, S["heading_box"], [
        {"runs": [{"t": o["eyebrow"].upper(), "size": S["eyebrow"]["size"], "bold": True, "colour": "green"}], "line_pts": S["eyebrow"]["line_pts"]},
        {"runs": [{"t": o["title_white"], "size": S["title"]["size"], "bold": True, "colour": "white", "tracking": -50},
                  {"t": o["title_green"], "size": S["title"]["size"], "bold": True, "colour": "green", "tracking": -50}],
         "line": S["title"]["line"]}], insets=(0.1, 0.05, 0.1, 0.05))
    cg, ct = S["cards"], S["card_text"]
    for k, c in enumerate(o["cards"]):
        x = cg["x"] + k * (cg["w"] + cg["gap"]); y = cg["y"]
        card(sl, {"x": x, "y": y, "w": cg["w"], "h": cg["h"]})
        text(sl, {"x": x + 0.22, "y": y + 0.3, "w": cg["w"] - 0.44, "h": ct["head_box_h"]}, [
            {"runs": [{"t": f"{k + 1:02d}", "size": ct["number_size"], "bold": True, "tracking": -50}], "line_pts": 40},
            {"runs": [{"t": c["head"], "size": ct["head_size"], "bold": True, "colour": "green", "tracking": -50}], "line": 0.8}],
            insets=(0.1, 0.05, 0.1, 0.05))
        rule(sl, x + 0.32, y + ct["rule_dy"], 0.6)
        text(sl, {"x": x + 0.22, "y": y + ct["body_dy"], "w": cg["w"] - 0.44, "h": cg["h"] - ct["body_dy"] - 0.2},
             [{"runs": [{"t": c["body"], "size": ct["body_size"]}]}], insets=(0.1, 0.05, 0.1, 0.05))
    footer(sl, len(prs.slides), year)
    notes(sl, o.get("notes"))

def closing(prs):
    S = B["slides"]["closing"]; sl = prs.slides.add_slide(BLANK)
    background(sl, S["background"])
    picture(sl, asset("closing_logo"), S["logo"])
    t = S["title"]
    text(sl, t, [{"runs": [{"t": t["text"], "font": "display", "size": t["size"]}], "line": t["line"]}], anchor=t["anchor"])

# ---------------------------------------------------------------- main
def bounded(value, limit, label):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError(f"{label} needs 1–{limit} characters")


def validate_outline(o, session_dir):
    _, expected, excluded = captures(session_dir)
    actual = [s for chapter in o["chapters"] for s in chapter["slides"]]
    if not expected or len(actual) != len(expected):
        raise ValueError("Outline must contain every eligible capture exactly once")
    for source, slide in zip(expected, actual):
        if any(slide.get(key) != source[key] for key in ("section_id", "image", "notes")):
            raise ValueError("Capture order, screenshot or exact narration changed; regenerate the outline")
        bounded(slide["headline"], 60, "Headline")
        if not 1 <= len(slide["takeaways"]) <= 3:
            raise ValueError("Use 1–3 takeaways; nothing is silently truncated")
        for takeaway in slide["takeaways"]:
            bounded(takeaway, 60, "Takeaway")
    for chapter in o["chapters"]:
        bounded(chapter.get("eyebrow", chapter["name"]), 24, "Chapter eyebrow")
        if o.get("dividers", False):
            bounded(chapter.get("divider_title", chapter["name"]), 20, "Divider title")
    if o.get("cover", False):
        bounded(o["title"]["lead"] + " " + o["title"].get("rest", ""), 60, "Cover title")
        if o.get("subtitle"):
            bounded(o["subtitle"], 45, "Subtitle")
    if o.get("summary"):
        if len(o["summary"]["cards"]) != 4:
            raise ValueError("Summary needs four supported cards, or omit it")
        for card in o["summary"]["cards"]:
            bounded(card["head"], 30, "Summary heading")
            bounded(card["body"], 110, "Summary body")
    for entry in excluded:
        print("Excluded:", entry["directory"], "-", entry["reason"])


def main(outline_path, session_dir, out):
    with open(outline_path, encoding="utf-8") as source:
        o = json.load(source)
    validate_outline(o, session_dir)
    year = o.get("year") or datetime.date.today().year
    global BLANK
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(B["canvas"]["width"]), Inches(B["canvas"]["height"])
    BLANK = prs.slide_layouts[6]
    if o.get("cover", False): cover(prs, o, year)
    i = 0
    for ci, ch in enumerate(o["chapters"], 1):
        if o.get("dividers", False):
            divider(prs, ci, ch.get("divider_title", ch["name"]), f"Part {ci}: {ch.get('divider_title', ch['name'])}.")
        for s in ch["slides"]:
            i += 1
            content(prs, i, ch.get("eyebrow", ch["name"]), s, session_dir, year)
    if o.get("summary"): summary(prs, o["summary"], year)
    if o.get("closing", False): closing(prs)
    prs.core_properties.title = (o["title"]["lead"] + " " + o["title"].get("rest", "")).strip()
    # Finish serialization before exclusively creating a new output file.
    data = io.BytesIO()
    prs.save(data)
    with new_output(out) as destination:
        destination.write(data.getvalue())
    print(f"saved {out}: {len(prs.slides)} slides, {i} capture slides")

if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
